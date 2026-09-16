import os
from pathlib import Path
import pypdf
import docx

class DocumentLoader:
    @staticmethod
    def extract_text(file_path: str) -> str:
        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")

        suffix = path.suffix.lower()
        if suffix == ".pdf":
            return DocumentLoader._extract_pdf(path)
        elif suffix in [".docx", ".doc"]:
            return DocumentLoader._extract_docx(path)
        elif suffix in [".txt", ".md", ".json", ".csv", ".py", ".js", ".html"]:
            return DocumentLoader._extract_plaintext(path)
        else:
            # Attempt plain text read as fallback
            try:
                return DocumentLoader._extract_plaintext(path)
            except Exception:
                raise ValueError(f"Unsupported file format: {suffix}")

    @staticmethod
    def _extract_pdf(path: Path) -> str:
        text_parts = []
        with open(path, "rb") as f:
            reader = pypdf.PdfReader(f)
            for page_idx, page in enumerate(reader.pages):
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text.strip())
        return "\n\n".join(text_parts)

    @staticmethod
    def _extract_docx(path: Path) -> str:
        doc = docx.Document(path)
        text_parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    text_parts.append(row_text)
        return "\n\n".join(text_parts)

    @staticmethod
    def _extract_plaintext(path: Path) -> str:
        for encoding in ["utf-8", "utf-8-sig", "latin-1", "cp1252"]:
            try:
                with open(path, "r", encoding=encoding) as f:
                    return f.read()
            except (UnicodeDecodeError, UnicodeError):
                continue
        raise ValueError(f"Could not decode text file: {path.name}")
